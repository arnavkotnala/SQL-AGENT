"""Excel (.xlsx) and CSV export of query results, plus a registry that serves files by opaque ID."""
from __future__ import annotations

import csv
import re
import threading
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

EXCEL_MAX_DATA_ROWS = 1_048_575  # 1,048,576 rows minus the header
EXCEL_MAX_CELL_CHARS = 32_767
_ISO_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_ISO_DATETIME = re.compile(r"^\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(:\d{2}(\.\d+)?)?$")
_PERCENT_NAME = re.compile(r"(percent|percentage|pct|ratio_pct|rate_pct|%)", re.I)
_CURRENCY_NAME = re.compile(r"(salary|fee|amount|price|revenue|cost|sales|income|total_sales)", re.I)
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

FORMATS = {
    "integer": "#,##0",
    "decimal": "#,##0.00",
    "currency": "#,##0.00",
    "percent_0_100": "0.00%",
    "percent_fraction": "0.00%",
    "date": "yyyy-mm-dd",
    "datetime": "yyyy-mm-dd hh:mm",
}


@dataclass
class ColumnSpec:
    name: str
    label: str
    kind: str  # text | integer | decimal | currency | percent_0_100 | percent_fraction | date | datetime


@dataclass
class ExportMetadata:
    request: str
    goal: str
    source: str
    sql: str
    params: dict[str, Any]
    truncated: bool
    row_limit: int
    summary: str | None = None


@dataclass
class ExportFile:
    id: str
    path: Path
    filename: str
    media_type: str
    rows: int
    created_at: str = field(default_factory=lambda: datetime.now().isoformat(timespec="seconds"))

    def to_dict(self) -> dict[str, Any]:
        return {"id": self.id, "filename": self.filename, "rows": self.rows, "created_at": self.created_at,
                "media_type": self.media_type, "url": f"/api/exports/{self.id}"}


def humanize(name: str) -> str:
    s = re.sub(r"[_\s]+", " ", str(name)).strip()
    s = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", s)
    words = []
    for w in s.split(" "):
        lw = w.lower()
        words.append(w.upper() if lw in {"id", "dob", "pct"} else (w if w.isupper() and len(w) > 1 else w.capitalize()))
    return " ".join(words) or str(name)


def infer_column_specs(columns: Sequence[str], rows: Sequence[Sequence[Any]],
                       hints: dict[str, tuple[str, str]] | None = None) -> list[ColumnSpec]:
    """Decide each column's Excel type from values (SQLite returns dates as text and derived columns are
    untyped), optionally refined by model-provided hints {column: (label, kind)}."""
    hints = {k.lower(): v for k, v in (hints or {}).items()}
    sample = rows[:5000]
    specs: list[ColumnSpec] = []
    for i, col in enumerate(columns):
        vals = [r[i] for r in sample if r[i] is not None and r[i] != ""]
        kind = "text"
        if vals:
            if all(isinstance(v, bool) for v in vals):
                kind = "text"
            elif all(isinstance(v, int) and not isinstance(v, bool) for v in vals):
                kind = "integer"
            elif all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
                kind = "decimal"
            elif all(isinstance(v, (datetime,)) or (isinstance(v, str) and _ISO_DATETIME.match(v)) for v in vals):
                kind = "datetime"
            elif all(isinstance(v, date) or (isinstance(v, str) and _ISO_DATE.match(v)) for v in vals):
                kind = "date"
        numeric = kind in ("integer", "decimal")
        if numeric and _PERCENT_NAME.search(col):
            kind = "percent_fraction" if all(0 <= float(v) <= 1 for v in vals) and not any(
                float(v) > 1 for v in vals) and _looks_fractional(vals) else "percent_0_100"
        elif kind == "decimal" and _CURRENCY_NAME.search(col):
            kind = "currency"
        label = humanize(col)
        hint = hints.get(col.lower())
        if hint:
            h_label, h_kind = hint
            label = h_label or label
            # Trust the model's numeric/date intent only when the values agree with it.
            if h_kind in ("percent_0_100", "percent_fraction", "currency", "decimal", "integer") and numeric:
                kind = h_kind if not (h_kind == "integer" and kind != "integer") else kind
            elif h_kind in ("date", "datetime") and kind in ("date", "datetime"):
                kind = h_kind
        specs.append(ColumnSpec(name=col, label=label, kind=kind))
    return specs


def _looks_fractional(vals: list[Any]) -> bool:
    return any(isinstance(v, float) and not float(v).is_integer() for v in vals)


def to_cell_value(v: Any, kind: str) -> Any:
    if v is None:
        return None
    if kind == "percent_0_100" and isinstance(v, (int, float)):
        return float(v) / 100.0
    if kind in ("date", "datetime") and isinstance(v, str):
        try:
            return datetime.fromisoformat(v.replace(" ", "T")) if kind == "datetime" else datetime.strptime(v, "%Y-%m-%d")
        except ValueError:
            return v
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if len(v) > EXCEL_MAX_CELL_CHARS:
            v = v[: EXCEL_MAX_CELL_CHARS - 15] + "...[truncated]"
    return v


def _text_len(v: Any) -> int:
    if v is None:
        return 0
    if isinstance(v, float):
        return len(f"{v:,.2f}")
    if isinstance(v, (datetime, date)):
        return 16 if isinstance(v, datetime) else 10
    return len(str(v))


HEADER_FILL = PatternFill("solid", start_color="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D7DE")
HEADER_BORDER = Border(bottom=Side(style="medium", color="1F3A5F"))


def write_xlsx(path: Path, columns: Sequence[str], rows: Sequence[Sequence[Any]], meta: ExportMetadata,
               hints: dict[str, tuple[str, str]] | None = None) -> int:
    """Write results to a genuine .xlsx file (streaming/write-only mode). Returns the data rows written."""
    specs = infer_column_specs(columns, rows, hints)
    truncated_for_excel = len(rows) > EXCEL_MAX_DATA_ROWS
    data = rows[:EXCEL_MAX_DATA_ROWS]

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Results")
    ws.freeze_panes = "A2"
    for i, spec in enumerate(specs, start=1):
        sample_w = max([_text_len(r[i - 1]) for r in data[:1000]] + [0])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(spec.label) + 2, sample_w + 2))
    header = []
    for spec in specs:
        c = WriteOnlyCell(ws, value=spec.label)
        c.font, c.fill, c.border = HEADER_FONT, HEADER_FILL, HEADER_BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        header.append(c)
    ws.append(header or [WriteOnlyCell(ws, value="(no columns)")])
    if columns and data:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(data) + 1}"

    if not data:
        msg = WriteOnlyCell(ws, value="No matching records were found for this request.")
        msg.font = Font(italic=True, color="7A7A7A")
        ws.append([msg])
    for row in data:
        out = []
        for v, spec in zip(row, specs):
            cv = to_cell_value(v, spec.kind)
            cell = WriteOnlyCell(ws, value=cv)
            if isinstance(cv, str) and cv.startswith(_FORMULA_PREFIXES):
                cell.data_type = "s"  # never let data be interpreted as a formula
            fmt = FORMATS.get(spec.kind)
            if fmt and cv is not None and not isinstance(cv, str):
                cell.number_format = fmt
            out.append(cell)
        ws.append(out)

    info = wb.create_sheet("Query Info")
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 100
    notes = []
    if meta.truncated:
        notes.append(f"Result limited to the first {meta.row_limit:,} rows; more rows matched. "
                     "Narrow the request or raise MAX_RESULT_ROWS to export everything.")
    if truncated_for_excel:
        notes.append(f"Excel allows at most {EXCEL_MAX_DATA_ROWS:,} data rows; use the CSV export for the rest.")
    meta_rows = [
        ("Request", meta.request),
        ("Interpreted goal", meta.goal),
        ("Data source", meta.source),
        ("Executed SQL", meta.sql),
        ("Parameters", ", ".join(f"{k} = {v!r}" for k, v in meta.params.items()) or "(none)"),
        ("Record count", len(data)),
        ("Complete result", "No - see notes" if notes else "Yes"),
        ("Notes", " ".join(notes) or "-"),
        ("Summary", meta.summary or "-"),
        ("Exported at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Generated by", "LoopSQL AI"),
    ]
    for k, v in meta_rows:
        kc = WriteOnlyCell(info, value=k)
        kc.font = Font(bold=True)
        vc = WriteOnlyCell(info, value=to_cell_value(v, "text") if isinstance(v, str) else v)
        if isinstance(v, str) and v.startswith(_FORMULA_PREFIXES):
            vc.data_type = "s"
        vc.alignment = Alignment(wrap_text=True, vertical="top")
        info.append([kc, vc])
    wb.save(path)
    return len(data)


def _csv_safe(v: Any) -> Any:
    if isinstance(v, str) and v.startswith(_FORMULA_PREFIXES):
        try:
            float(v)
            return v
        except ValueError:
            return "'" + v
    return v


def write_csv(path: Path, columns: Sequence[str], rows: Iterable[Sequence[Any]]) -> int:
    n = 0
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(columns)
        batch = []
        for r in rows:
            batch.append([_csv_safe(v) for v in r])
            n += 1
            if len(batch) >= 1000:
                w.writerows(batch)
                batch.clear()
        w.writerows(batch)
    return n


def verify_xlsx(path: Path, columns: Sequence[str], expected_rows: int) -> None:
    """Re-open the workbook and check header + row count. Raises AssertionError on mismatch."""
    wb = load_workbook(path, read_only=True)
    try:
        ws = wb["Results"]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise AssertionError("Results sheet is empty.")
        if len([c for c in rows[0] if c is not None]) != max(1, len(columns)):
            raise AssertionError("Header does not match the result columns.")
        data_rows = rows[1:]
        actual = 0 if (expected_rows == 0) else len(data_rows)
        if expected_rows == 0 and len(data_rows) != 1:
            raise AssertionError("Expected an explanatory 'no records' row.")
        if actual != expected_rows:
            raise AssertionError(f"Workbook has {actual} data rows, expected {expected_rows}.")
    finally:
        wb.close()


def slugify(text: str, max_len: int = 40) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", (text or "").lower()).strip("_")
    return (s[:max_len].rstrip("_") or "results")


class ExportStore:
    """Files are saved under exports_dir/<random id>/ and served only by that ID."""

    def __init__(self, exports_dir: Path):
        self.exports_dir = exports_dir
        self._files: dict[str, ExportFile] = {}
        self._lock = threading.Lock()

    def new_path(self, filename: str) -> tuple[str, Path]:
        export_id = uuid.uuid4().hex
        d = self.exports_dir / export_id
        d.mkdir(parents=True, exist_ok=True)
        return export_id, d / filename

    def register(self, export_id: str, path: Path, media_type: str, rows: int) -> ExportFile:
        ef = ExportFile(id=export_id, path=path, filename=path.name, media_type=media_type, rows=rows)
        with self._lock:
            self._files[export_id] = ef
        return ef

    def register_existing(self, path: Path, media_type: str, rows: int) -> ExportFile:
        """Register a file already written inside exports_dir (e.g. an edited workbook copy)."""
        resolved = path.resolve()
        if self.exports_dir.resolve() not in resolved.parents:
            raise ValueError("Refusing to serve a file outside the exports directory.")
        return self.register(uuid.uuid4().hex, resolved, media_type, rows)

    def get(self, export_id: str) -> ExportFile | None:
        if not re.fullmatch(r"[0-9a-f]{32}", export_id or ""):
            return None
        with self._lock:
            ef = self._files.get(export_id)
        if ef is None or not ef.path.is_file():
            return None
        return ef


XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLSM_MEDIA = "application/vnd.ms-excel.sheet.macroEnabled.12"
CSV_MEDIA = "text/csv"
