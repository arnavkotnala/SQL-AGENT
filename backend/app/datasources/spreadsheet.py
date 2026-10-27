"""Excel (.xlsx/.xlsm) and CSV files as data sources.

Each sheet is imported into a private SQLite *mirror* so the agent can query it with SQL. The SQLite
rowid of every record equals its Excel row number, which lets an approved modification be mapped back to
exact cells. Modifications are never applied directly: they are dry-run on the mirror inside a
transaction that is rolled back, the cell-level diff is shown for approval, and only then the workbook is
written (to a new copy by default, or in place after a backup).
"""
from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import shutil
import sqlite3
import threading
import time
import uuid
import zipfile
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime, time as dtime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from .base import DataSourceError, QueryExecutionError, SchemaInfo
from .sql_adapter import SQLiteAdapter, _SQLITE_DENIED_FUNCTIONS

EXCEL_SUFFIXES = {".xlsx", ".xlsm"}
CSV_SUFFIXES = {".csv"}
SPREADSHEET_SUFFIXES = EXCEL_SUFFIXES | CSV_SUFFIXES

# Identifiers that would need quoting in SQL; suffixing them keeps generated SQL simple and valid.
_RESERVED = {
    "select", "from", "where", "group", "order", "by", "limit", "offset", "table", "index", "join", "on",
    "and", "or", "not", "null", "is", "in", "as", "case", "when", "then", "else", "end", "values", "insert",
    "update", "delete", "create", "drop", "alter", "into", "set", "union", "all", "distinct", "having",
    "primary", "key", "default", "check", "references", "foreign", "unique", "between", "like", "exists",
    "to", "with", "cross", "inner", "outer", "left", "right", "full", "natural", "using", "desc", "asc",
    "date", "time", "timestamp", "user", "column", "row", "rows", "transaction", "commit", "rollback",
    "rowid", "oid", "_rowid_",
}


class StaleSourceError(Exception):
    """The file changed on disk after the change preview was produced."""


@dataclass
class SheetColumn:
    name: str  # SQL identifier
    header: str  # original header text
    col_idx: int  # 1-based Excel column index
    sql_type: str  # INTEGER | REAL | TEXT | DATE | DATETIME | BOOLEAN
    is_percent: bool = False
    has_formula: bool = False


@dataclass
class SheetTable:
    table: str
    sheet: str
    header_row: int
    columns: list[SheetColumn]
    data_rows: list[int]  # Excel row numbers that hold records
    sheet_max_col: int = 0  # right-most used column anywhere on the sheet

    @property
    def last_row(self) -> int:
        return max(self.data_rows) if self.data_rows else self.header_row

    def column(self, name: str) -> SheetColumn:
        for c in self.columns:
            if c.name.lower() == name.lower():
                return c
        raise KeyError(name)


@dataclass
class CellChange:
    row: int
    column: str
    header: str
    col_letter: str
    old: Any
    new: Any


@dataclass
class ChangeSet:
    id: str
    table: str
    sheet: str
    statement_type: str
    sql: str
    params: dict[str, Any]
    updates: list[CellChange] = field(default_factory=list)
    inserted_rows: list[dict[str, Any]] = field(default_factory=list)
    deleted_rows: list[dict[str, Any]] = field(default_factory=list)  # includes "_excel_row"
    new_columns: list[SheetColumn] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    fingerprint: str = ""

    @property
    def affected_rows(self) -> int:
        return len({u.row for u in self.updates}) + len(self.inserted_rows) + len(self.deleted_rows)

    def to_dict(self, max_items: int = 500) -> dict[str, Any]:
        return {
            "id": self.id,
            "table": self.table,
            "sheet": self.sheet,
            "statement_type": self.statement_type,
            "sql": self.sql,
            "affected_rows": self.affected_rows,
            "updated_cells": len(self.updates),
            "new_columns": [{"name": c.name, "header": c.header, "col_letter": get_column_letter(c.col_idx)}
                            for c in self.new_columns],
            "updates": [u.__dict__ for u in self.updates[:max_items]],
            "inserted_rows": self.inserted_rows[:max_items],
            "deleted_rows": self.deleted_rows[:max_items],
            "truncated_preview": max(len(self.updates), len(self.inserted_rows), len(self.deleted_rows)) > max_items,
            "warnings": self.warnings,
        }


# --------------------------------------------------------------------------------------------
# Value helpers
# --------------------------------------------------------------------------------------------
def sanitize_identifier(raw: Any, fallback: str) -> str:
    s = str(raw if raw is not None else "").strip().lower()
    s = s.replace("%", " pct ").replace("#", " no ").replace("&", " and ")
    s = re.sub(r"[^a-z0-9_]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    if not s:
        s = fallback
    if s[0].isdigit():
        s = "c_" + s
    if s in _RESERVED:
        s = s + "_"
    return s[:60]


def _dedupe(name: str, used: set[str]) -> str:
    base, n = name, 2
    while name in used:
        name = f"{base}_{n}"
        n += 1
    used.add(name)
    return name


def _to_sql_value(v: Any) -> Any:
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d") if v.time() == dtime(0) else v.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, dtime):
        return v.strftime("%H:%M:%S")
    if isinstance(v, str):
        return v.strip() if v.strip() else None
    return v


def _infer_type(values: list[Any]) -> str:
    vals = [v for v in values if v is not None and not (isinstance(v, str) and not v.strip())]
    if not vals:
        return "TEXT"
    if all(isinstance(v, bool) for v in vals):
        return "BOOLEAN"
    if all(isinstance(v, int) and not isinstance(v, bool) for v in vals):
        return "INTEGER"
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals):
        return "REAL"
    if all(isinstance(v, (datetime, date)) for v in vals):
        if all(not isinstance(v, datetime) or v.time() == dtime(0) for v in vals):
            return "DATE"
        return "DATETIME"
    return "TEXT"


def _coerce_for_type(v: Any, sql_type: str) -> Any:
    v = _to_sql_value(v)
    if v is None:
        return None
    if sql_type == "TEXT" and not isinstance(v, str):
        if isinstance(v, float) and v.is_integer():
            return str(int(v))
        return str(v)
    return v


def to_excel_value(v: Any, col: SheetColumn) -> Any:
    """Convert a mirror value back into a native Excel cell value."""
    if v is None:
        return None
    if col.sql_type in ("DATE", "DATETIME") and isinstance(v, str):
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(v, fmt)
            except ValueError:
                continue
        return v
    if col.sql_type == "BOOLEAN" and isinstance(v, (int, float)):
        return bool(v)
    if col.sql_type in ("INTEGER", "REAL") and isinstance(v, str):
        try:
            f = float(v)
            return int(f) if f.is_integer() and col.sql_type == "INTEGER" else f
        except ValueError:
            return v
    return v


def file_fingerprint(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


# --------------------------------------------------------------------------------------------
# Import
# --------------------------------------------------------------------------------------------
def _pick_header_row(rows: list[tuple[int, list[Any]]]) -> int | None:
    """First row (within the first 20) whose filled-cell count is at least half of the widest row."""
    head = rows[:20]
    if not head:
        return None
    counts = [(r, sum(1 for v in vals if v not in (None, ""))) for r, vals in head]
    widest = max(c for _, c in counts)
    if widest == 0:
        return None
    for r, c in counts:
        if c >= max(1, widest / 2):
            return r
    return None


def _read_excel_rows(path: Path) -> tuple[dict[str, list[tuple[int, list[Any]]]], dict[str, dict[tuple[int, int], bool]], dict[str, dict[tuple[int, int], str]]]:
    """Returns per-sheet values, formula map and number formats (read-only/streaming mode)."""
    try:
        wb_v = load_workbook(path, data_only=True, read_only=True)
        wb_f = load_workbook(path, data_only=False, read_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError) as e:
        raise DataSourceError(f"Could not open Excel file '{path.name}': {e}") from None
    values: dict[str, list[tuple[int, list[Any]]]] = {}
    formulas: dict[str, dict[tuple[int, int], bool]] = {}
    formats: dict[str, dict[tuple[int, int], str]] = {}
    try:
        for ws in wb_v.worksheets:
            rows: list[tuple[int, list[Any]]] = []
            fmts: dict[tuple[int, int], str] = {}
            for r_idx, row in enumerate(ws.iter_rows(), start=1):
                vals = []
                for c_idx, cell in enumerate(row, start=1):
                    vals.append(getattr(cell, "value", None))
                    nf = getattr(cell, "number_format", None)
                    if nf and "%" in nf:
                        fmts[(r_idx, c_idx)] = nf
                rows.append((r_idx, vals))
            values[ws.title] = rows
            formats[ws.title] = fmts
        for ws in wb_f.worksheets:
            fm: dict[tuple[int, int], bool] = {}
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                for c_idx, v in enumerate(row, start=1):
                    if isinstance(v, str) and v.startswith("="):
                        fm[(r_idx, c_idx)] = True
            formulas[ws.title] = fm
    finally:
        wb_v.close()
        wb_f.close()
    return values, formulas, formats


def _read_csv_rows(path: Path) -> dict[str, list[tuple[int, list[Any]]]]:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [(i, [_csv_cell(v) for v in r]) for i, r in enumerate(csv.reader(io.StringIO(text), dialect), start=1)]
    return {path.stem: rows}


_INT_RE = re.compile(r"^[+-]?\d+$")
_FLOAT_RE = re.compile(r"^[+-]?(\d+\.\d*|\.\d+|\d+)([eE][+-]?\d+)?$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")


def _csv_cell(v: str) -> Any:
    s = v.strip()
    if s == "":
        return None
    if _INT_RE.match(s) and not (len(s) > 1 and s.lstrip("+-").startswith("0")):
        return int(s)
    if _FLOAT_RE.match(s):
        return float(s)
    if _DATE_RE.match(s):
        try:
            return datetime.strptime(s, "%Y-%m-%d").date()
        except ValueError:
            return s
    return s


def import_spreadsheet(path: Path, mirror_path: Path) -> list[SheetTable]:
    suffix = path.suffix.lower()
    if suffix in EXCEL_SUFFIXES:
        values, formulas, formats = _read_excel_rows(path)
    elif suffix in CSV_SUFFIXES:
        values, formulas, formats = _read_csv_rows(path), {}, {}
    else:
        raise DataSourceError(f"Unsupported spreadsheet type '{suffix}'. Use .xlsx, .xlsm or .csv.")

    tables: list[SheetTable] = []
    used_tables: set[str] = set()
    tmp = mirror_path.with_suffix(f".{uuid.uuid4().hex[:8]}.tmp")
    conn = sqlite3.connect(tmp)
    try:
        for sheet, rows in values.items():
            header_row = _pick_header_row(rows)
            if header_row is None:
                continue
            header_vals = dict(rows)[header_row]
            filled = [i for i, v in enumerate(header_vals, start=1) if v not in (None, "")]
            if not filled:
                continue
            first_col, last_col = filled[0], filled[-1]
            data = [(r, vals) for r, vals in rows if r > header_row]
            # widen to columns that hold data even without a header
            for _, vals in data:
                for i, v in enumerate(vals, start=1):
                    if v not in (None, "") and i > last_col:
                        last_col = i
            used_cols: set[str] = set()
            columns: list[SheetColumn] = []
            sheet_formulas = formulas.get(sheet, {})
            sheet_formats = formats.get(sheet, {})
            for ci in range(first_col, last_col + 1):
                header = header_vals[ci - 1] if ci - 1 < len(header_vals) else None
                letter = get_column_letter(ci)
                name = _dedupe(sanitize_identifier(header, f"column_{letter.lower()}"), used_cols)
                col_vals = [vals[ci - 1] if ci - 1 < len(vals) else None for _, vals in data]
                numeric_cells = [(r, ci) for (r, vals), v in zip(data, col_vals) if isinstance(v, (int, float)) and not isinstance(v, bool)]
                pct = bool(numeric_cells) and sum(1 for k in numeric_cells if k in sheet_formats) >= len(numeric_cells) / 2
                columns.append(SheetColumn(
                    name=name,
                    header=str(header) if header not in (None, "") else f"Column {letter}",
                    col_idx=ci,
                    sql_type=_infer_type(col_vals),
                    is_percent=pct,
                    has_formula=any((r, ci) in sheet_formulas for r, _ in data),
                ))
            table = _dedupe(sanitize_identifier(sheet, "sheet"), used_tables)
            data_rows: list[int] = []
            col_defs = ", ".join(f'"{c.name}" {c.sql_type}' for c in columns)
            conn.execute(f'CREATE TABLE "{table}" ({col_defs})')
            placeholders = ", ".join("?" for _ in range(len(columns) + 1))
            col_list = ", ".join(f'"{c.name}"' for c in columns)
            batch = []
            for r, vals in data:
                row_vals = [vals[c.col_idx - 1] if c.col_idx - 1 < len(vals) else None for c in columns]
                if all(v in (None, "") or (isinstance(v, str) and not v.strip()) for v in row_vals):
                    continue
                data_rows.append(r)
                batch.append([r] + [_coerce_for_type(v, c.sql_type) for v, c in zip(row_vals, columns)])
            conn.executemany(f'INSERT INTO "{table}" (rowid, {col_list}) VALUES ({placeholders})', batch)
            sheet_max_col = max((max((i for i, v in enumerate(vals, start=1) if v not in (None, "")), default=0)
                                 for _, vals in rows), default=0)
            tables.append(SheetTable(table=table, sheet=sheet, header_row=header_row, columns=columns,
                                     data_rows=data_rows, sheet_max_col=sheet_max_col))
        conn.commit()
    finally:
        conn.close()
    if not tables:
        tmp.unlink(missing_ok=True)
        raise DataSourceError(f"No tabular data found in '{path.name}'. The first rows should contain column headers.")
    _replace_with_retry(tmp, mirror_path)
    return tables


def _replace_with_retry(src: Path, dst: Path, attempts: int = 10) -> None:
    for i in range(attempts):
        try:
            os.replace(src, dst)
            return
        except PermissionError:
            if i == attempts - 1:
                raise
            time.sleep(0.1)


def workbook_feature_warnings(path: Path) -> list[str]:
    """Features openpyxl may not round-trip when the workbook is re-saved."""
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        return []
    try:
        with zipfile.ZipFile(path) as z:
            names = z.namelist()
    except zipfile.BadZipFile:
        return []
    found = []
    checks = {"xl/charts/": "charts", "xl/drawings/": "drawings/shapes", "xl/media/": "images",
              "xl/pivotTables/": "pivot tables", "xl/slicers/": "slicers"}
    for prefix, label in checks.items():
        if any(n.startswith(prefix) for n in names):
            found.append(label)
    if not found:
        return []
    return [f"This workbook contains {', '.join(found)}. The spreadsheet library used to write changes may not "
            "preserve them, so saving to a new copy (not overwriting) is strongly recommended."]


# --------------------------------------------------------------------------------------------
# The data source
# --------------------------------------------------------------------------------------------
class SpreadsheetSource(SQLiteAdapter):
    def __init__(self, path: Path, workspace_dir: Path, display_name: str | None = None,
                 origin: str = "path"):
        self.file_path = Path(path).resolve()
        if not self.file_path.is_file():
            raise DataSourceError(f"File not found: {self.file_path}")
        self.origin = origin  # "path" (user's own file on disk) | "upload"
        self.lock = threading.RLock()
        self.mirror_path = workspace_dir / f"{uuid.uuid4().hex}.sqlite"
        self.tables: list[SheetTable] = []
        self._stat: tuple[int, int] | None = None
        self._reimport()
        super().__init__(self.mirror_path, display_name or self.file_path.name)
        self.display_target = str(self.file_path) if origin == "path" else self.file_path.name

    @property
    def editable(self) -> bool:  # type: ignore[override]
        return self.file_path.suffix.lower() in EXCEL_SUFFIXES

    @property
    def kind(self) -> str:
        return "excel" if self.file_path.suffix.lower() in EXCEL_SUFFIXES else "csv"

    def _reimport(self) -> None:
        with self.lock:
            st = self.file_path.stat()
            self.tables = import_spreadsheet(self.file_path, self.mirror_path)
            self._stat = (st.st_mtime_ns, st.st_size)

    def refresh_if_changed(self) -> bool:
        with self.lock:
            try:
                st = self.file_path.stat()
            except OSError:
                raise DataSourceError(f"File is no longer available: {self.file_path.name}") from None
            if self._stat == (st.st_mtime_ns, st.st_size):
                return False
            self._reimport()
            return True

    def sheet_table(self, table: str) -> SheetTable:
        for t in self.tables:
            if t.table.lower() == table.lower():
                return t
        raise KeyError(table)

    def describe(self) -> dict[str, Any]:
        d = super().describe()
        d.update({"kind": self.kind, "file": self.file_path.name, "origin": self.origin,
                  "sheets": [{"table": t.table, "sheet": t.sheet, "rows": len(t.data_rows)} for t in self.tables]})
        return d

    def _table_kind(self, kind: str) -> str:
        return "sheet"

    def _table_note(self, table: str) -> str | None:
        try:
            t = self.sheet_table(table)
        except KeyError:
            return None
        return f"from sheet '{t.sheet}' (header on row {t.header_row})"

    def _column_note(self, table: str, column: str) -> str | None:
        try:
            c = self.sheet_table(table).column(column)
        except KeyError:
            return None
        notes = []
        if c.header != c.name:
            notes.append(f"header '{c.header}'")
        if c.is_percent:
            notes.append("Excel percentage stored as a fraction: 0.45 means 45%")
        if c.sql_type in ("DATE", "DATETIME"):
            notes.append("ISO text YYYY-MM-DD" + (" HH:MM:SS" if c.sql_type == "DATETIME" else ""))
        if c.has_formula:
            notes.append("calculated by a formula (read-only)")
        return "; ".join(notes) or None

    def inspect_schema(self) -> SchemaInfo:
        with self.lock:
            return super().inspect_schema()

    def execute_read(self, sql, params, max_rows, timeout_s):  # type: ignore[override]
        with self.lock:
            return super().execute_read(sql, params, max_rows, timeout_s)

    def read_only_columns(self) -> dict[str, set[str]]:
        return {t.table: {c.name for c in t.columns if c.has_formula} for t in self.tables}

    # ---- modifications ---------------------------------------------------------------------
    def plan_new_columns(self, target_table: str, new_columns: list[dict[str, Any]] | None) -> list[SheetColumn]:
        """Validate requested new columns and assign them the next free spreadsheet columns."""
        if not new_columns:
            return []
        if len(new_columns) > 5:
            raise QueryExecutionError("At most 5 new columns can be added in one change.", recoverable=True)
        sheet = self.sheet_table(target_table)
        used = {c.name for c in sheet.columns}
        next_idx = max(sheet.sheet_max_col, max(c.col_idx for c in sheet.columns)) + 1
        planned: list[SheetColumn] = []
        for nc in new_columns:
            name = sanitize_identifier(nc.get("name") or nc.get("header"), "new_column")
            if name in used:
                continue  # already exists: the statement can simply use it
            used.add(name)
            header = re.sub(r"\s+", " ", str(nc.get("header") or humanize_header(name))).strip()[:100]
            sql_type = str(nc.get("type") or "TEXT").upper()
            if sql_type not in ("TEXT", "INTEGER", "REAL", "DATE"):
                sql_type = "TEXT"
            planned.append(SheetColumn(name=name, header=header, col_idx=next_idx, sql_type=sql_type))
            next_idx += 1
        return planned

    def dry_run(self, sql: str, params: dict[str, Any], target_table: str, max_rows: int,
                timeout_s: float, new_columns: list[SheetColumn] | None = None) -> ChangeSet:
        if not self.editable:
            raise QueryExecutionError("CSV sources are read-only; only Excel workbooks can be modified.",
                                      recoverable=False)
        with self.lock:
            sheet = self.sheet_table(target_table)
            new_columns = new_columns or []
            all_columns = sheet.columns + new_columns
            cols = [c.name for c in all_columns]
            col_sql = ", ".join(f'"{c}"' for c in cols)
            old_sql = ", ".join(f'"{c.name}"' for c in sheet.columns)
            conn = sqlite3.connect(self.mirror_path, isolation_level=None, check_same_thread=False)
            try:
                pad = (None,) * len(new_columns)
                before = {r[0]: tuple(r[1:]) + pad
                          for r in conn.execute(f'SELECT rowid, {old_sql} FROM "{sheet.table}"')}
                deadline = time.monotonic() + timeout_s
                conn.execute("BEGIN")
                for nc in new_columns:  # identifiers were sanitized by plan_new_columns
                    conn.execute(f'ALTER TABLE "{sheet.table}" ADD COLUMN "{nc.name}" {nc.sql_type}')
                conn.set_authorizer(_write_authorizer(sheet.table))
                conn.set_progress_handler(lambda: 1 if time.monotonic() > deadline else 0, 2_000)
                try:
                    conn.execute(sql, params)
                except sqlite3.Error as e:
                    msg = str(e)
                    if "interrupted" in msg:
                        msg = f"Statement exceeded the {timeout_s:g}s timeout."
                    raise QueryExecutionError(msg, recoverable="not authorized" not in msg) from None
                finally:
                    conn.set_authorizer(None)
                    conn.set_progress_handler(None, 0)
                after = {r[0]: r[1:] for r in conn.execute(f'SELECT rowid, {col_sql} FROM "{sheet.table}"')}
                conn.execute("ROLLBACK")
            finally:
                conn.close()

        cs = ChangeSet(id=uuid.uuid4().hex, table=sheet.table, sheet=sheet.sheet,
                       statement_type=sql.split(None, 1)[0].lower(), sql=sql, params=params,
                       fingerprint=file_fingerprint(self.file_path), new_columns=list(new_columns))
        for rowid, old in before.items():
            if rowid not in after:
                rec = {c: v for c, v in zip(cols, old)}
                rec["_excel_row"] = rowid
                cs.deleted_rows.append(rec)
                continue
            new = after[rowid]
            for c, o, n in zip(all_columns, old, new):
                if o != n:
                    cs.updates.append(CellChange(row=rowid, column=c.name, header=c.header,
                                                 col_letter=get_column_letter(c.col_idx), old=o, new=n))
        for rowid, new in after.items():
            if rowid not in before:
                cs.inserted_rows.append({c: v for c, v in zip(cols, new)})
        if new_columns and not cs.updates and not cs.inserted_rows:
            cs.warnings.append("The new column(s) would be added but no values are filled in.")
        if cs.affected_rows > max_rows:
            raise QueryExecutionError(
                f"The modification would affect {cs.affected_rows} rows, above the limit of {max_rows}. "
                "Narrow the request.", recoverable=False)
        if cs.deleted_rows:
            cs.warnings.append("Deleting rows shifts the rows below them up. Row-by-row formulas are adjusted, but "
                               "formulas that reference ranges (e.g. a SUM over the column) are not rewritten.")
        cs.warnings.extend(workbook_feature_warnings(self.file_path))
        return cs

    def apply(self, cs: ChangeSet, mode: str, exports_dir: Path, backups_dir: Path) -> dict[str, Any]:
        """Write an approved ChangeSet. mode: 'copy' (new file) or 'in_place' (backup, then overwrite)."""
        if mode not in ("copy", "in_place"):
            raise ValueError("mode must be 'copy' or 'in_place'")
        with self.lock:
            if file_fingerprint(self.file_path) != cs.fingerprint:
                raise StaleSourceError("The file changed on disk after the preview was created. "
                                       "Run the request again to get a fresh preview.")
            sheet = self.sheet_table(cs.table)
            keep_vba = self.file_path.suffix.lower() == ".xlsm"
            wb = load_workbook(self.file_path, keep_vba=keep_vba)
            ws = wb[sheet.sheet]
            expected: dict[tuple[int, int], Any] = {}
            columns_by_name = {c.name.lower(): c for c in sheet.columns + cs.new_columns}
            header_template = ws.cell(row=sheet.header_row, column=sheet.columns[-1].col_idx)
            for nc in cs.new_columns:
                hc = ws.cell(row=sheet.header_row, column=nc.col_idx)
                if header_template.has_style:
                    hc._style = copy(header_template._style)
                _set_value(hc, nc.header)
                expected[(sheet.header_row, nc.col_idx)] = nc.header
                ws.column_dimensions[get_column_letter(nc.col_idx)].width = max(12, len(nc.header) + 4)
            for u in cs.updates:
                col = columns_by_name[u.column.lower()]
                val = to_excel_value(u.new, col)
                _set_value(ws.cell(row=u.row, column=col.col_idx), val)
                expected[(u.row, col.col_idx)] = val
            if cs.deleted_rows:
                _delete_rows_keeping_formulas(ws, sorted(d["_excel_row"] for d in cs.deleted_rows))
            if cs.inserted_rows:
                start = sheet.last_row + 1
                template_row = sheet.last_row if sheet.data_rows else None
                if any(c.value is not None for r in ws.iter_rows(min_row=start, max_row=start + len(cs.inserted_rows) - 1)
                       for c in r):
                    ws.insert_rows(start, amount=len(cs.inserted_rows))
                for i, rec in enumerate(cs.inserted_rows):
                    r = start + i
                    for col in sheet.columns + cs.new_columns:
                        cell = ws.cell(row=r, column=col.col_idx)
                        tmpl = ws.cell(row=template_row, column=col.col_idx) if template_row else None
                        if tmpl is not None and tmpl.has_style:
                            cell._style = copy(tmpl._style)
                        if col.has_formula and tmpl is not None and isinstance(tmpl.value, str) and tmpl.value.startswith("="):
                            # continue the table's calculated column into the new row
                            cell.value = Translator(tmpl.value, origin=tmpl.coordinate).translate_formula(cell.coordinate)
                            continue
                        val = to_excel_value(rec.get(col.name), col)
                        _set_value(cell, val)
                        expected[(r, col.col_idx)] = val

            if mode == "copy":
                out_dir = exports_dir / uuid.uuid4().hex
                out_dir.mkdir(parents=True, exist_ok=True)
                out_path = out_dir / f"{self.file_path.stem}_edited{self.file_path.suffix.lower()}"
                wb.save(out_path)
                backup_path = None
            else:
                backups_dir.mkdir(parents=True, exist_ok=True)
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                backup_path = backups_dir / f"{self.file_path.stem}_backup_{stamp}{self.file_path.suffix.lower()}"
                shutil.copy2(self.file_path, backup_path)
                tmp = self.file_path.with_name(f".~loopsql_{uuid.uuid4().hex[:8]}{self.file_path.suffix}")
                wb.save(tmp)
                try:
                    _replace_with_retry(tmp, self.file_path)
                except PermissionError:
                    tmp.unlink(missing_ok=True)
                    raise PermissionError(
                        f"Could not overwrite '{self.file_path.name}'. Close it in Excel (or any other program) "
                        "and approve again. Your original file was not changed.") from None
                out_path = self.file_path
            wb.close()

            mismatches = _verify_written(out_path, sheet.sheet, expected)
            if mismatches:
                raise RuntimeError(f"Verification failed for {len(mismatches)} cell(s), e.g. {mismatches[0]}.")
            if mode == "in_place":
                self._reimport()
            return {"mode": mode, "output_path": str(out_path), "output_name": out_path.name,
                    "backup_path": str(backup_path) if backup_path else None,
                    "verified_cells": len(expected), "affected_rows": cs.affected_rows}

    def close(self) -> None:
        super().close()
        try:
            self.mirror_path.unlink(missing_ok=True)
        except OSError:
            pass


def _delete_rows_keeping_formulas(ws: Any, rows: list[int]) -> None:
    """Delete rows bottom-up, then re-translate formulas in rows that moved up so row-relative references
    (e.g. =E5*2 on row 5) keep pointing at their own row. Cross-range references are not rewritten."""
    first = rows[0]
    deleted = set(rows)
    formulas: list[tuple[int, int, str]] = []
    for r in ws.iter_rows(min_row=first + 1, max_row=ws.max_row):
        for c in r:
            if isinstance(c.value, str) and c.value.startswith("=") and c.row not in deleted:
                formulas.append((c.row, c.column, c.value))
    for row in sorted(rows, reverse=True):
        ws.delete_rows(row)
    for orig_row, col, formula in formulas:
        shift = sum(1 for d in rows if d < orig_row)
        if shift:
            src = f"{get_column_letter(col)}{orig_row}"
            dst = f"{get_column_letter(col)}{orig_row - shift}"
            ws.cell(row=orig_row - shift, column=col).value = Translator(formula, origin=src).translate_formula(dst)


def humanize_header(name: str) -> str:
    return " ".join(w.capitalize() for w in name.split("_") if w) or name


def _set_value(cell: Any, val: Any) -> None:
    """Assign a value without letting text such as '=HYPERLINK(...)' become a live formula."""
    cell.value = val
    if isinstance(val, str) and val.startswith("="):
        cell.data_type = "s"


def _verify_written(path: Path, sheet: str, expected: dict[tuple[int, int], Any]) -> list[str]:
    """Re-open the saved workbook and compare every written cell."""
    wb = load_workbook(path, read_only=False)
    try:
        ws = wb[sheet]
        bad = []
        for (r, c), val in expected.items():
            got = ws.cell(row=r, column=c).value
            if not _same(got, val):
                bad.append(f"{get_column_letter(c)}{r}: expected {val!r}, found {got!r}")
        return bad
    finally:
        wb.close()


def _same(a: Any, b: Any) -> bool:
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool):
        return abs(float(a) - float(b)) < 1e-9
    return a == b


def _write_authorizer(target_table: str):
    target = target_table.lower()
    allowed_reads = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_RECURSIVE}

    def auth(action: int, arg1: str | None, arg2: str | None, _db: str | None, _trig: str | None) -> int:
        if action in allowed_reads:
            return sqlite3.SQLITE_OK
        if action == sqlite3.SQLITE_FUNCTION:
            return sqlite3.SQLITE_DENY if (arg2 or "").lower() in _SQLITE_DENIED_FUNCTIONS else sqlite3.SQLITE_OK
        if action in (sqlite3.SQLITE_UPDATE, sqlite3.SQLITE_INSERT, sqlite3.SQLITE_DELETE):
            return sqlite3.SQLITE_OK if (arg1 or "").lower() == target else sqlite3.SQLITE_DENY
        return sqlite3.SQLITE_DENY

    return auth
