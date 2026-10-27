"""Shared fixtures. The language model is replaced by a scripted fake so tests are deterministic;
see test_gemini_live.py for the optional real-Gemini integration test."""
from __future__ import annotations

import threading
import time
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from app.agent.loop import Agent
from app.agent.run import AgentRun, RunOptions, RunStatus
from app.config import Settings
from app.datasources.registry import SourceRegistry
from app.export.exporter import ExportStore
from app.llm.schemas import GoalInterpretation, ResultEvaluation, SqlProposal, SqlParam, NewColumn
from app.sample_data.generate import generate_all


class FakeLLM:
    """Scripted stand-in for Gemini. `script[task]` is a list consumed in order; each item is a model
    instance, an Exception (raised), or a callable(prompt) -> model."""

    name = "fake-llm"

    def __init__(self, script: dict[str, list[Any]] | None = None, default_eval: bool = True):
        self.script = {k: list(v) for k, v in (script or {}).items()}
        self.calls: list[dict[str, str]] = []
        self.default_eval = default_eval

    def generate_structured(self, *, task, system, prompt, schema):
        self.calls.append({"task": task, "system": system, "prompt": prompt})
        queue = self.script.get(task)
        if not queue:
            if task in ("evaluate", "evaluate_change") and self.default_eval:
                return ok_eval()
            raise AssertionError(f"FakeLLM has no scripted response for task '{task}'")
        item = queue.pop(0) if len(queue) > 1 else queue[0]
        if isinstance(item, Exception):
            raise item
        if callable(item) and not isinstance(item, (GoalInterpretation, SqlProposal, ResultEvaluation)):
            item = item(prompt)
        assert isinstance(item, schema), f"scripted {type(item).__name__} for {schema.__name__}"
        return item

    def prompts(self, task: str) -> list[str]:
        return [c["prompt"] for c in self.calls if c["task"] == task]


def goal(operation: str = "read", feasible: bool = True, reason: str | None = None,
         clarify: str | None = None, text: str = "Answer the user's question") -> GoalInterpretation:
    return GoalInterpretation(interpreted_goal=text, operation=operation, feasible=feasible,
                              infeasible_reason=reason, needs_clarification=bool(clarify),
                              clarification_question=clarify, success_criteria=["correct rows"],
                              requested_columns=[], relevant_tables=[])


def proposal(sql: str, params: dict[str, Any] | None = None, new_columns: list[dict] | None = None) -> SqlProposal:
    return SqlProposal(plan=["step"], sql=sql, rationale="test",
                       params=[SqlParam(name=k, value=v) for k, v in (params or {}).items()],
                       new_columns=[NewColumn(**c) for c in (new_columns or [])])


def ok_eval(summary: str = "Looks right.", formats: list[dict] | None = None) -> ResultEvaluation:
    return ResultEvaluation(satisfied=True, summary=summary, column_formats=formats or [])


def bad_eval(instructions: str) -> ResultEvaluation:
    return ResultEvaluation(satisfied=False, issues=["wrong"], refinement_instructions=instructions, summary="no")


@pytest.fixture(scope="session")
def sample_dir(tmp_path_factory) -> Path:
    d = tmp_path_factory.mktemp("samples")
    generate_all(d)
    return d


@pytest.fixture
def settings(tmp_path, sample_dir) -> Settings:
    s = Settings(_env_file=None, data_dir=tmp_path / "data", gemini_api_key=None, approval_timeout_seconds=20,
                 frontend_dist=tmp_path / "no_dist", gemini_max_retries=2,
                 trusted_hosts=["127.0.0.1", "localhost", "testserver"])
    s.ensure_dirs()
    for f in sample_dir.iterdir():  # reuse pre-generated samples (read-only)
        (s.samples_dir / f.name).write_bytes(f.read_bytes())
    return s


class Harness:
    def __init__(self, settings: Settings, llm: FakeLLM):
        self.settings = settings
        self.llm = llm
        self.registry = SourceRegistry(settings)
        self.exports = ExportStore(settings.exports_dir)
        self.agent = Agent(settings, self.registry, self.exports, lambda: self.llm)

    def source(self, sample: str = "education") -> str:
        return self.registry.add_sample(sample).id

    def run(self, request: str, source_id: str | None = None, **opts) -> AgentRun:
        run = AgentRun(request=request, options=RunOptions(source_id=source_id, **opts),
                       max_iterations=self.settings.agent_max_iterations)
        self.agent.run(run)
        return run

    def run_async(self, request: str, source_id: str | None = None, **opts) -> tuple[AgentRun, threading.Thread]:
        run = AgentRun(request=request, options=RunOptions(source_id=source_id, **opts),
                       max_iterations=self.settings.agent_max_iterations)
        t = threading.Thread(target=self.agent.run, args=(run,), daemon=True)
        t.start()
        return run, t


def wait_until(pred: Callable[[], bool], timeout: float = 15.0) -> None:
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        if pred():
            return
        time.sleep(0.02)
    raise AssertionError("condition not reached in time")


def wait_pending(run: AgentRun, kind: str) -> dict:
    wait_until(lambda: (run.pending or {}).get("type") == kind or run.status in RunStatus.TERMINAL)
    assert (run.pending or {}).get("type") == kind, f"run ended with {run.status}: {run.errors}"
    return run.pending


@pytest.fixture
def make_harness(settings) -> Callable[[FakeLLM], Harness]:
    def _make(llm: FakeLLM) -> Harness:
        return Harness(settings, llm)
    return _make


def make_student_workbook(path: Path) -> Path:
    """A realistic user workbook: title row, bold header, percent formatting, dates, formula column."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Students"
    ws.append(["Semester 5 attendance"])
    ws.append(["Roll No", "Student Name", "Attendance %", "Last Updated", "Credits", "Score x2"])
    for c in ws[2]:
        c.font = Font(bold=True)
    data = [("GE001", "Asha Rawat", 0.45, 18), ("GE002", "Ravi Negi", 0.82, 20), ("GE003", "Kiran Bisht", 0.31, 16),
            ("GE004", "Meera Joshi", 0.67, 22), ("GE005", "Dev Thapa", 0.49, 18)]
    for i, (roll, name, pct, credits) in enumerate(data, start=3):
        ws.append([roll, name, pct, datetime(2026, 10, 1), credits, f"=E{i}*2"])
        ws.cell(row=i, column=3).number_format = "0%"
    other = wb.create_sheet("Notes")
    other.append(["Key", "Value"])
    other.append(["Owner", "Dept office"])
    wb.save(path)
    return path
