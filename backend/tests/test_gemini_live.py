"""Optional integration test against the real Gemini API.

Runs only when GEMINI_API_KEY is configured (environment or backend/.env):
    .venv\\Scripts\\python -m pytest -m live -s
"""
import sqlite3

import pytest

from app.agent.loop import Agent
from app.agent.run import AgentRun, RunOptions, RunStatus
from app.config import Settings
from app.datasources.registry import SourceRegistry
from app.export.exporter import ExportStore
from app.llm.provider import build_provider

pytestmark = pytest.mark.live

_real = Settings()
if not _real.gemini_configured:
    pytest.skip("GEMINI_API_KEY not configured; live Gemini test skipped.", allow_module_level=True)


@pytest.fixture
def live_agent(settings):
    settings.gemini_api_key = _real.gemini_api_key
    settings.gemini_model = _real.gemini_model
    registry = SourceRegistry(settings)
    agent = Agent(settings, registry, ExportStore(settings.exports_dir), lambda: build_provider(settings))
    return agent, registry, settings


@pytest.mark.parametrize("request_text", [
    "Find all students whose attendance is below 50% and export their names, roll numbers, and attendance percentages.",
    "Calculate average marks for each course.",
])
def test_live_education_queries(live_agent, request_text):
    agent, registry, settings = live_agent
    run = AgentRun(request=request_text, options=RunOptions(source_id=registry.add_sample("education").id),
                   max_iterations=settings.agent_max_iterations)
    agent.run(run)
    print("\nSQL:", run.sql, "\nrows:", len(run.rows), "iterations:", run.iteration, "\nsummary:", run.summary)
    assert run.status == RunStatus.COMPLETED, run.errors
    assert run.rows and run.exports
    if "average" in request_text:
        assert len(run.rows) == 5  # one row per course


def _run_to_approval(agent, run):
    import threading
    t = threading.Thread(target=agent.run, args=(run,), daemon=True)
    t.start()
    for _ in range(1800):  # up to ~3 minutes of model calls
        if run.pending or run.status in RunStatus.TERMINAL:
            break
        t.join(0.1)
    return t


def test_live_excel_query(live_agent, tmp_path):
    from .conftest import make_student_workbook
    agent, registry, settings = live_agent
    src = registry.add_local_path(str(make_student_workbook(tmp_path / "class.xlsx")))
    run = AgentRun(request="Which students have attendance below 50%? Show roll number, name and attendance.",
                   options=RunOptions(source_id=src.id), max_iterations=5)
    agent.run(run)
    print("\nSQL:", run.sql, "\nrows:", run.rows)
    assert run.status == RunStatus.COMPLETED, run.errors
    roll_idx = next(i for i, col in enumerate(run.columns) if "roll" in col.lower())
    assert sorted(r[roll_idx] for r in run.rows) == ["GE001", "GE003", "GE005"]  # 45%, 31%, 49%


def test_live_excel_modification_in_place(live_agent, tmp_path):
    """The headline workflow: plain-English instruction -> Gemini plans an Excel edit -> preview ->
    approval -> the user's workbook is changed on disk (with backup) and verified."""
    from openpyxl import load_workbook
    from .conftest import make_student_workbook
    agent, registry, settings = live_agent
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    src = registry.add_local_path(str(xlsx))
    run = AgentRun(request="Add a new column called Status: write 'Detained' for students whose attendance is "
                           "below 50%, and 'Eligible' for everyone else.",
                   options=RunOptions(source_id=src.id, write_mode="in_place"), max_iterations=5)
    t = _run_to_approval(agent, run)
    print("\nSQL:", run.sql, "\nnew columns:", (run.change_preview or {}).get("new_columns"),
          "\nerrors:", run.errors)
    assert (run.pending or {}).get("type") == "change_approval", (run.status, run.errors)
    run.respond("change_approval", {"approved": True, "mode": "in_place"})
    t.join(60)
    assert run.status == RunStatus.COMPLETED, run.errors
    wb = load_workbook(xlsx)
    ws = wb["Students"]
    headers = [c.value for c in ws[2]]
    status_col = next(i for i, h in enumerate(headers) if h and "status" in str(h).lower())
    got = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=status_col + 1).value for r in range(3, 8)}
    assert got == {"GE001": "Detained", "GE002": "Eligible", "GE003": "Detained", "GE004": "Eligible",
                   "GE005": "Detained"}
    assert ws["F3"].value == "=E3*2"  # formulas untouched
    assert list(settings.backups_dir.iterdir())  # backup written first


def test_live_attendance_matches_ground_truth(live_agent):
    agent, registry, settings = live_agent
    run = AgentRun(request="List the roll numbers of students whose overall attendance across all subjects is below 50%.",
                   options=RunOptions(source_id=registry.add_sample("education").id), max_iterations=5)
    agent.run(run)
    assert run.status == RunStatus.COMPLETED, run.errors
    with sqlite3.connect(settings.samples_dir / "education.db") as c:
        truth = {r[0] for r in c.execute(
            "SELECT s.roll_number FROM students s JOIN attendance a USING (student_id) GROUP BY s.student_id "
            "HAVING SUM(a.classes_attended) * 100.0 / SUM(a.classes_held) < 50")}
    roll_idx = next(i for i, col in enumerate(run.columns) if "roll" in col.lower())
    assert {r[roll_idx] for r in run.rows} == truth
