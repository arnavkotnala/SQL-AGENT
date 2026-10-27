"""Agent loop on read-only databases: success paths, correction loop, iteration limits, domain independence."""
import sqlite3
from datetime import date, timedelta

from openpyxl import load_workbook

from app.agent.run import RunStatus
from app.llm.provider import LLMNotConfiguredError, LLMRateLimitError

from .conftest import FakeLLM, bad_eval, goal, ok_eval, proposal

ATTENDANCE_SQL = """
SELECT s.first_name || ' ' || s.last_name AS student_name, s.roll_number,
       ROUND(SUM(a.classes_attended) * 100.0 / SUM(a.classes_held), 2) AS attendance_percentage
FROM students s JOIN attendance a ON a.student_id = s.student_id
GROUP BY s.student_id, s.first_name, s.last_name, s.roll_number
HAVING SUM(a.classes_attended) * 100.0 / SUM(a.classes_held) < :threshold
ORDER BY attendance_percentage, s.roll_number
"""

AVG_SQL = """
SELECT c.course_name, ROUND(AVG(m.marks_obtained), 2) AS average_marks
FROM marks m JOIN students s ON s.student_id = m.student_id JOIN courses c ON c.course_id = s.course_id
GROUP BY c.course_id, c.course_name ORDER BY c.course_name
"""


def _db(h):
    return sqlite3.connect(h.settings.samples_dir / "education.db")


def _xlsx_rows(h, run):
    ef = h.exports.get(run.exports[0]["id"])
    wb = load_workbook(ef.path, read_only=True)
    rows = list(wb["Results"].iter_rows(values_only=True))
    info = {r[0]: r[1] for r in wb["Query Info"].iter_rows(values_only=True)}
    wb.close()
    return rows, info


def test_nl_to_sql_success_and_attendance_below_50(make_harness):
    llm = FakeLLM({"interpret": [goal(text="Students with attendance below 50%")],
                   "generate_sql": [proposal(ATTENDANCE_SQL, {"threshold": 50})],
                   "evaluate": [ok_eval(formats=[{"column": "attendance_percentage", "label": "Attendance %",
                                                   "kind": "percent_0_100"}])]})
    h = make_harness(llm)
    run = h.run("Find all students whose attendance is below 50%.", h.source())
    assert run.status == RunStatus.COMPLETED, run.errors
    assert run.iteration == 1 and run.export_status == "done"

    # independent ground truth straight from the database
    with _db(h) as c:
        truth = c.execute("""SELECT s.roll_number, SUM(a.classes_attended)*100.0/SUM(a.classes_held) p
                             FROM students s JOIN attendance a USING (student_id) GROUP BY s.student_id
                             HAVING p < 50 ORDER BY round(p, 2), s.roll_number""").fetchall()
    assert len(truth) > 5
    assert [r[1] for r in run.rows] == [t[0] for t in truth]
    assert all(r[2] < 50 for r in run.rows)

    # the schema was given to the model (grounding) and the request reached it
    assert "TABLE attendance" in llm.prompts("generate_sql")[0]
    assert "attendance is below 50%" in llm.prompts("interpret")[0]

    # Excel export contains exactly the database records, with % stored as a real percentage
    rows, info = _xlsx_rows(h, run)
    assert rows[0] == ("Student Name", "Roll Number", "Attendance %")
    assert len(rows) - 1 == len(truth)
    for xrow, (roll, pct) in zip(rows[1:], truth):
        assert xrow[1] == roll
        assert abs(xrow[2] - round(pct, 2) / 100) < 1e-9
    assert info["Record count"] == len(truth)
    assert "HAVING" in info["Executed SQL"]


def test_joins_and_aggregation_average_marks_per_course(make_harness):
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal(AVG_SQL)]}))
    run = h.run("Calculate average marks for each course.", h.source())
    assert run.status == RunStatus.COMPLETED, run.errors
    with _db(h) as c:
        truth = c.execute("""SELECT c.course_name, ROUND(AVG(m.marks_obtained), 2) FROM marks m
                             JOIN students s USING (student_id) JOIN courses c USING (course_id)
                             GROUP BY c.course_id ORDER BY c.course_name""").fetchall()
    assert [tuple(r) for r in run.rows] == truth and len(truth) == 5


def test_scored_below_40_in_any_subject(make_harness):
    sql = """SELECT DISTINCT s.roll_number, s.first_name, s.last_name, sub.subject_name, m.marks_obtained
             FROM marks m JOIN students s ON s.student_id = m.student_id
             JOIN subjects sub ON sub.subject_id = m.subject_id WHERE m.marks_obtained < :min_marks
             ORDER BY s.roll_number, sub.subject_name"""
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal(sql, {"min_marks": 40})]}))
    run = h.run("Find students who scored below 40 in any subject.", h.source())
    assert run.status == RunStatus.COMPLETED
    assert run.rows and all(r[4] < 40 for r in run.rows)


def test_empty_result_is_exported_with_message(make_harness):
    sql = "SELECT roll_number, first_name FROM students WHERE enrollment_year = :year"
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal(sql, {"year": 2015})],
                              "evaluate": [ok_eval("No students enrolled in 2015.")]}))
    run = h.run("Find students enrolled in 2015.", h.source())
    assert run.status == RunStatus.COMPLETED
    assert run.rows == [] and run.result_meta["row_count"] == 0
    rows, info = _xlsx_rows(h, run)
    assert rows[0] == ("Roll Number", "First Name")
    assert "No matching records" in rows[1][0]
    assert info["Record count"] == 0


def test_invalid_column_is_corrected_on_next_iteration(make_harness):
    llm = FakeLLM({"interpret": [goal()], "generate_sql": [
        proposal("SELECT roll_no, attendance FROM students"),  # hallucinated identifiers
        proposal("SELECT roll_number FROM students WHERE enrollment_year = :y", {"y": 2024}),
    ]})
    h = make_harness(llm)
    run = h.run("List roll numbers of 2024 students", h.source())
    assert run.status == RunStatus.COMPLETED, run.errors
    assert run.iteration == 2
    assert run.attempts[0]["outcome"] == "validation_error"
    # the error was fed back to the model for the correction
    assert "roll_no" in llm.prompts("generate_sql")[1] and "could not be resolved" in llm.prompts("generate_sql")[1]


def test_invalid_table_and_syntax_error_then_success(make_harness):
    llm = FakeLLM({"interpret": [goal()], "generate_sql": [
        proposal("SELECT * FROM pupils"),
        proposal("SELEC roll_number FROM students"),
        proposal("SELECT roll_number FROM students"),
    ]})
    h = make_harness(llm)
    run = h.run("List all roll numbers", h.source())
    assert run.status == RunStatus.COMPLETED
    assert [a["outcome"] for a in run.attempts] == ["validation_error", "validation_error", "success"]
    assert "Unknown table 'pupils'" in run.attempts[0]["error"]
    assert "syntax" in run.attempts[1]["error"].lower()
    assert len(run.rows) == 150


def test_database_execution_error_is_recovered(make_harness):
    llm = FakeLLM({"interpret": [goal()], "generate_sql": [
        proposal("SELECT no_such_function(first_name) AS x FROM students"),
        proposal("SELECT UPPER(first_name) AS x FROM students"),
    ]})
    h = make_harness(llm)
    run = h.run("Uppercase names", h.source())
    assert run.status == RunStatus.COMPLETED
    assert run.attempts[0]["outcome"] == "execution_error"
    assert "no such function" in run.attempts[0]["error"].lower()


def test_unsatisfied_evaluation_triggers_refinement(make_harness):
    llm = FakeLLM({"interpret": [goal()],
                   "generate_sql": [proposal("SELECT roll_number FROM students"),
                                    proposal("SELECT roll_number, email FROM students")],
                   "evaluate": [bad_eval("Include the email column."), ok_eval()]})
    h = make_harness(llm)
    run = h.run("Roll numbers and emails", h.source())
    assert run.status == RunStatus.COMPLETED and run.iteration == 2
    assert "Include the email column." in llm.prompts("generate_sql")[1]


def test_repeated_failures_stop_at_iteration_limit(make_harness, settings):
    bad = [proposal(f"SELECT missing_col_{i} FROM students") for i in range(10)]
    llm = FakeLLM({"interpret": [goal()], "generate_sql": bad})
    h = make_harness(llm)
    run = h.run("Something impossible", h.source())
    assert run.status == RunStatus.FAILED
    assert run.iteration == settings.agent_max_iterations == 5
    assert len(llm.prompts("generate_sql")) == 5
    assert "within 5 iterations" in run.errors[-1]
    assert run.exports == []


def test_identical_queries_are_not_reexecuted_and_loop_stops(make_harness):
    same = proposal("SELECT roll_number FROM students")
    llm = FakeLLM({"interpret": [goal()], "generate_sql": [same, same, same, same]},)
    llm.script["evaluate"] = [bad_eval("Wrong."), bad_eval("Wrong.")]
    h = make_harness(llm)
    run = h.run("x", h.source())
    assert run.status == RunStatus.FAILED
    outcomes = [a["outcome"] for a in run.attempts]
    assert outcomes == ["unsatisfied", "duplicate", "duplicate"]
    assert "same query" in run.errors[-1]
    assert len(llm.prompts("evaluate")) == 1  # executed exactly once


def test_infeasible_request_needs_input(make_harness):
    llm = FakeLLM({"interpret": [goal(feasible=False, reason="There is no employee or salary data.")]})
    h = make_harness(llm)
    run = h.run("Find employees earning more than the average salary in their department.", h.source())
    assert run.status == RunStatus.NEEDS_INPUT
    assert "salary" in run.summary
    assert llm.prompts("generate_sql") == []


def test_modify_request_on_database_is_blocked(make_harness):
    h = make_harness(FakeLLM({"interpret": [goal(operation="modify")]}))
    run = h.run("Delete all students with low attendance", h.source())
    assert run.status == RunStatus.BLOCKED
    assert "read-only" in run.errors[-1]


def test_same_agent_works_on_hospital_domain(make_harness):
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    sql = """SELECT a.appointment_id, d.name AS department, a.appointment_date
             FROM appointments a JOIN departments d ON d.department_id = a.department_id
             WHERE a.appointment_date = :tomorrow ORDER BY a.appointment_id"""
    llm = FakeLLM({"interpret": [goal()], "generate_sql": [proposal(sql, {"tomorrow": tomorrow})]})
    h = make_harness(llm)
    run = h.run("Find hospital appointments scheduled for tomorrow and export appointment IDs, departments, and dates.",
                h.source("hospital"))
    assert run.status == RunStatus.COMPLETED
    assert len(run.rows) == 12 and all(r[2] == tomorrow for r in run.rows)
    assert "TABLE appointments" in llm.prompts("generate_sql")[0]
    assert date.today().isoformat() in llm.calls[1]["system"]  # today's date given for relative dates


def test_gemini_not_configured_fails_clearly(make_harness):
    h = make_harness(FakeLLM())

    def factory():
        raise LLMNotConfiguredError("Gemini API key is not configured.")
    h.agent.llm_factory = factory
    run = h.run("anything", h.source())
    assert run.status == RunStatus.FAILED and "not configured" in run.errors[-1]


def test_gemini_rate_limit_during_run_fails_with_message(make_harness):
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [LLMRateLimitError("Gemini rate limit (429)")]}))
    run = h.run("anything", h.source())
    assert run.status == RunStatus.FAILED and "429" in run.errors[-1]


def test_result_size_limit_is_disclosed(make_harness, settings):
    settings.max_result_rows = 10
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal("SELECT * FROM marks")]}))
    run = h.run("all marks", h.source())
    assert run.status == RunStatus.COMPLETED
    assert len(run.rows) == 10 and run.result_meta["truncated"] is True
    rows, info = _xlsx_rows(h, run)
    assert len(rows) == 11
    assert info["Complete result"].startswith("No") and "first 10 rows" in info["Notes"]
