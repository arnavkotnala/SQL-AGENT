"""HTTP API end-to-end with the fake model: run lifecycle, downloads, limits, auth, SPA serving."""
import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import create_app

from .conftest import FakeLLM, bad_eval, goal, make_student_workbook, proposal

SQL = "SELECT roll_number, first_name FROM students WHERE course_id = :cid ORDER BY roll_number"


def _client(settings, llm):
    return TestClient(create_app(settings, llm_factory=lambda: llm))


def _wait(c, run_id, statuses=("completed", "failed", "blocked", "needs_input", "cancelled", "awaiting_approval")):
    for _ in range(300):
        r = c.get(f"/api/runs/{run_id}").json()
        if r["status"] in statuses:
            return r
        time.sleep(0.03)
    raise AssertionError(r)


def test_full_run_and_excel_download(settings):
    llm = FakeLLM({"interpret": [goal()], "generate_sql": [proposal(SQL, {"cid": 2})]})
    with _client(settings, llm) as c:
        sid = c.post("/api/sources", json={"type": "sample", "value": "education"}).json()["id"]
        run = c.post("/api/runs", json={"request": "Students of course 2", "source_id": sid}).json()
        final = _wait(c, run["id"])
        assert final["status"] == "completed", final["errors"]
        assert [e["phase"] for e in final["events"]][:3] == ["Observing", "Inspecting schema", "Inspecting schema"]
        assert final["events"][-1]["phase"] == "Completed"
        xlsx = next(e for e in final["exports"] if e["format"] == "xlsx")
        resp = c.get(xlsx["url"])
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert "attachment" in resp.headers["content-disposition"]
        wb = load_workbook(io.BytesIO(resp.content))
        rows = list(wb["Results"].iter_rows(values_only=True))
        api_rows = c.get(f"/api/runs/{run['id']}/rows", params={"page_size": 500}).json()
        assert [list(r) for r in rows[1:]] == api_rows["rows"] and api_rows["total"] == 30
        csv = c.get(next(e for e in final["exports"] if e["format"] == "csv")["url"])
        assert csv.status_code == 200 and csv.text.lstrip("\ufeff").startswith("roll_number,first_name")
        # pagination + search
        page = c.get(f"/api/runs/{run['id']}/rows", params={"page": 2, "page_size": 10}).json()
        assert len(page["rows"]) == 10 and page["total"] == 30
        first_name = api_rows["rows"][0][1]
        found = c.get(f"/api/runs/{run['id']}/rows", params={"search": first_name.lower()}).json()
        assert found["total"] >= 1 and all(first_name in r[1] for r in found["rows"])
        assert c.get("/api/runs").json()[0]["id"] == run["id"]


@pytest.mark.parametrize("bad_id", ["nope", "0" * 32, "..%2F..%2Fapp%2Fconfig.py", "../../app/config.py"])
def test_download_endpoint_rejects_unknown_and_traversal(settings, bad_id):
    with _client(settings, FakeLLM()) as c:
        assert c.get(f"/api/exports/{bad_id}").status_code == 404


def test_upload_and_modify_via_api_with_approval(settings, tmp_path):
    data = make_student_workbook(tmp_path / "class.xlsx").read_bytes()
    llm = FakeLLM({"interpret": [goal(operation="modify")],
                   "generate_sql": [proposal("UPDATE students SET credits = 0 WHERE roll_no = 'GE001'")]})
    with _client(settings, llm) as c:
        up = c.post("/api/sources/upload", files={"file": ("class.xlsx", data)}).json()
        assert up["kind"] == "excel" and up["editable"]
        run = c.post("/api/runs", json={"request": "set GE001 credits to 0", "source_id": up["id"]}).json()
        waiting = _wait(c, run["id"])
        assert waiting["status"] == "awaiting_approval" and waiting["pending"]["type"] == "change_approval"
        assert waiting["pending"]["can_in_place"] is False  # uploads have no original on the server's disk
        assert c.post(f"/api/runs/{run['id']}/respond", json={"type": "sql_approval", "approved": True}).status_code == 409
        c.post(f"/api/runs/{run['id']}/respond", json={"type": "change_approval", "approved": True, "mode": "browser"})
        final = _wait(c, run["id"], ("completed", "failed"))
        assert final["status"] == "completed", final["errors"]
        assert final["apply_result"]["write_back"] == "browser" and "output_path" not in final["apply_result"]
        edited = c.get(final["exports"][0]["url"])
        wb = load_workbook(io.BytesIO(edited.content))
        assert wb["Students"]["E3"].value == 0


def test_upload_size_limit_and_bad_connections(settings, tmp_path):
    settings.max_upload_mb = 0
    with _client(settings, FakeLLM()) as c:
        r = c.post("/api/sources/upload", files={"file": ("a.xlsx", b"x" * 10)})
        assert r.status_code == 413
        r = c.post("/api/sources", json={"type": "path", "value": "C:/nope/missing.xlsx"})
        assert r.status_code == 400 and "not found" in r.json()["detail"]
        r = c.post("/api/sources", json={"type": "url", "value": "postgresql+psycopg://u:TopSecret@localhost:1/x"})
        assert r.status_code == 400 and "TopSecret" not in r.text
        assert c.get("/api/sources/doesnotexist/schema").status_code == 404


def test_partial_export_when_run_failed(settings):
    llm = FakeLLM({"interpret": [goal()], "generate_sql": [proposal("SELECT roll_number FROM students")],
                   "evaluate": [bad_eval("no")]})
    settings.agent_max_iterations = 1
    with _client(settings, llm) as c:
        sid = c.post("/api/sources", json={"type": "sample", "value": "education"}).json()["id"]
        run = c.post("/api/runs", json={"request": "all roll numbers", "source_id": sid}).json()
        assert _wait(c, run["id"])["status"] == "failed"
        ef = c.post(f"/api/runs/{run['id']}/export", params={"fmt": "xlsx"}).json()
        wb = load_workbook(io.BytesIO(c.get(ef["url"]).content))
        info = {r[0]: r[1] for r in wb["Query Info"].iter_rows(values_only=True)}
        assert "not confirmed" in info["Summary"]


def test_access_token_gate(settings):
    settings.app_access_token = "s3cret-key"
    with _client(settings, FakeLLM()) as c:
        assert c.get("/api/health").status_code == 200 and c.get("/api/health").json()["auth_required"]
        assert c.get("/api/sources").status_code == 401
        assert c.get("/api/sources", headers={"Authorization": "Bearer wrong"}).status_code == 401
        assert c.get("/api/sources", headers={"Authorization": "Bearer s3cret-key"}).status_code == 200


def test_health_never_exposes_the_api_key(settings):
    settings.gemini_api_key = "AIza-very-secret"
    with _client(settings, FakeLLM()) as c:
        r = c.get("/api/health")
        assert r.json()["gemini_configured"] is True and "AIza" not in r.text


def test_built_frontend_is_served(settings, tmp_path):
    dist = tmp_path / "dist"
    (dist / "assets").mkdir(parents=True)
    (dist / "index.html").write_text("<html>LoopSQL</html>")
    (dist / "assets" / "app.js").write_text("console.log(1)")
    settings.frontend_dist = dist
    with _client(settings, FakeLLM()) as c:
        assert "LoopSQL" in c.get("/").text
        assert "LoopSQL" in c.get("/some/client/route").text
        assert c.get("/assets/app.js").text == "console.log(1)"
        assert "LoopSQL" in c.get("/..%2F..%2Fapp%2Fconfig.py").text  # falls back to index, never leaks files
        assert c.get("/api/unknown").status_code == 404


def test_startup_purges_stale_files_but_keeps_backups(settings):
    import os
    old_export = settings.exports_dir / "a"
    old_export.mkdir()
    (old_export / "x.xlsx").write_bytes(b"x")
    os.utime(old_export, (0, 0))
    fresh_export = settings.exports_dir / "b"
    fresh_export.mkdir()
    (settings.workspaces_dir / "m.sqlite").write_bytes(b"x")
    (settings.backups_dir / "keep.xlsx").write_bytes(b"x")
    with _client(settings, FakeLLM()):
        pass
    assert not old_export.exists() and fresh_export.exists()
    assert list(settings.workspaces_dir.iterdir()) == []
    assert (settings.backups_dir / "keep.xlsx").exists()


def test_local_mode_rejects_foreign_host_headers(settings):
    with _client(settings, FakeLLM()) as c:
        assert c.get("/api/health", headers={"host": "evil.example"}).status_code == 400
        assert c.get("/api/health", headers={"host": "127.0.0.1:8000"}).status_code == 200
