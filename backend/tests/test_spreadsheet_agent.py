"""User Excel files: querying, approved modifications (copy / in place), and the agent asking for input."""
from openpyxl import load_workbook

from app.agent.run import RunStatus

from .conftest import FakeLLM, bad_eval, goal, make_student_workbook, ok_eval, proposal, wait_pending, wait_until

MARK_SQL = ("UPDATE students SET status = CASE WHEN attendance_pct < :limit THEN 'Detained' ELSE 'Eligible' END")
NEW_STATUS = [{"name": "status", "header": "Status", "type": "TEXT"}]


def _cells(path, sheet="Students"):
    wb = load_workbook(path)
    try:
        return [[c.value for c in r] for r in wb[sheet].iter_rows()]
    finally:
        wb.close()


def test_query_excel_file_and_export(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    sql = "SELECT roll_no, student_name, attendance_pct FROM students WHERE attendance_pct < :limit ORDER BY roll_no"
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal(sql, {"limit": 0.5})]}))
    src = h.registry.add_local_path(str(xlsx)).id
    run = h.run("Students below 50% attendance", src)
    assert run.status == RunStatus.COMPLETED, run.errors
    assert [r[0] for r in run.rows] == ["GE001", "GE003", "GE005"]
    schema_prompt = h.llm.prompts("interpret")[0]
    assert "0.45 means 45%" in schema_prompt and "read-only" in schema_prompt  # percent + formula notes


def test_modify_excel_copy_mode_requires_approval(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    original = xlsx.read_bytes()
    llm = FakeLLM({"interpret": [goal(operation="modify")],
                   "generate_sql": [proposal(MARK_SQL, {"limit": 0.5}, NEW_STATUS)]})
    h = make_harness(llm)
    src = h.registry.add_local_path(str(xlsx)).id
    run, t = h.run_async("Add a Status column: Detained if attendance below 50%, else Eligible", src)
    pending = wait_pending(run, "change_approval")
    change = pending["change"]
    assert change["affected_rows"] == 5 and change["new_columns"][0]["col_letter"] == "G"
    assert {u["new"] for u in change["updates"]} == {"Detained", "Eligible"}
    assert xlsx.read_bytes() == original  # nothing written before approval

    run.respond("change_approval", {"approved": True, "mode": "copy"})
    t.join(10)
    assert run.status == RunStatus.COMPLETED, run.errors
    assert xlsx.read_bytes() == original  # copy mode leaves the user's file untouched
    edited = h.exports.get(run.exports[0]["id"]).path
    cells = _cells(edited)
    assert cells[1][6] == "Status"
    assert [r[6] for r in cells[2:7]] == ["Detained", "Eligible", "Detained", "Eligible", "Detained"]
    assert cells[2][5] == "=E3*2"  # formulas preserved


def test_modify_excel_in_place_with_backup(make_harness, tmp_path, settings):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    original = xlsx.read_bytes()
    sql = "UPDATE students SET credits = credits + :bonus WHERE attendance_pct >= :limit"
    h = make_harness(FakeLLM({"interpret": [goal(operation="modify")],
                              "generate_sql": [proposal(sql, {"bonus": 2, "limit": 0.6})]}))
    src = h.registry.add_local_path(str(xlsx)).id
    run, t = h.run_async("Give 2 bonus credits to students with at least 60% attendance", src, write_mode="in_place")
    pending = wait_pending(run, "change_approval")
    assert pending["can_in_place"] is True and pending["default_mode"] == "in_place"
    run.respond("change_approval", {"approved": True, "mode": "in_place"})
    t.join(10)
    assert run.status == RunStatus.COMPLETED, run.errors
    cells = _cells(xlsx)
    assert [r[4] for r in cells[2:7]] == [18, 22, 16, 24, 18]
    backups = list(settings.backups_dir.iterdir())
    assert len(backups) == 1 and backups[0].read_bytes() == original
    # the mirror was refreshed: a follow-up query sees the new values
    adapter = h.registry.get(src).adapter
    assert adapter.execute_read("SELECT SUM(credits) FROM students", {}, 10, 5).rows == [(98,)]


def test_rejected_change_writes_nothing(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    original = xlsx.read_bytes()
    h = make_harness(FakeLLM({"interpret": [goal(operation="modify")],
                              "generate_sql": [proposal("DELETE FROM students WHERE attendance_pct < 0.5")]}))
    run, t = h.run_async("Remove students below 50%", h.registry.add_local_path(str(xlsx)).id)
    pending = wait_pending(run, "change_approval")
    assert len(pending["change"]["deleted_rows"]) == 3 and pending["change"]["warnings"]
    run.respond("change_approval", {"approved": False})
    t.join(10)
    assert run.status == RunStatus.CANCELLED and xlsx.read_bytes() == original


def test_formula_column_protected_then_corrected(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    llm = FakeLLM({"interpret": [goal(operation="modify")],
                   "generate_sql": [proposal("UPDATE students SET score_x2 = 0"),
                                    proposal("UPDATE students SET credits = 0 WHERE roll_no = :r", {"r": "GE002"})],
                   "evaluate_change": [ok_eval()]})
    h = make_harness(llm)
    run, t = h.run_async("Reset GE002's credits", h.registry.add_local_path(str(xlsx)).id)
    wait_pending(run, "change_approval")
    assert "formulas" in run.attempts[0]["error"]
    run.respond("change_approval", {"approved": True})
    t.join(10)
    assert run.status == RunStatus.COMPLETED


def test_wrong_change_is_refined_before_approval(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    llm = FakeLLM({"interpret": [goal(operation="modify")],
                   "generate_sql": [proposal("UPDATE students SET credits = 0"),
                                    proposal("UPDATE students SET credits = 0 WHERE roll_no = 'GE003'")],
                   "evaluate_change": [bad_eval("Only GE003 should change."), ok_eval()]})
    h = make_harness(llm)
    run, t = h.run_async("Set GE003 credits to 0", h.registry.add_local_path(str(xlsx)).id)
    pending = wait_pending(run, "change_approval")
    assert pending["change"]["affected_rows"] == 1 and run.iteration == 2
    run.cancel()
    t.join(10)


def test_csv_cannot_be_modified(make_harness, tmp_path):
    p = tmp_path / "x.csv"
    p.write_text("a,b\n1,2\n", encoding="utf-8")
    h = make_harness(FakeLLM({"interpret": [goal(operation="modify")]}))
    run = h.run("change b to 3", h.registry.add_local_path(str(p)).id)
    assert run.status == RunStatus.BLOCKED and "CSV" in run.errors[-1]


def test_agent_asks_for_a_file_when_none_is_connected(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal("SELECT roll_no FROM students")]}))
    run, t = h.run_async("List all roll numbers in my sheet")
    pending = wait_pending(run, "source")
    assert "Excel file" in pending["message"] and run.status == RunStatus.AWAITING_INPUT
    sid = h.registry.add_local_path(str(xlsx)).id
    run.respond("source", {"source_id": sid})
    t.join(10)
    assert run.status == RunStatus.COMPLETED and len(run.rows) == 5


def test_file_path_in_request_is_connected_automatically(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "my class.xlsx")
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal("SELECT COUNT(*) AS n FROM students")]}))
    run = h.run(f'How many students are in "{xlsx}"?')
    assert run.status == RunStatus.COMPLETED, run.errors
    assert run.rows == [(5,)] and run.source["kind"] == "excel"


def test_cloud_link_in_request_explains_and_asks(make_harness):
    h = make_harness(FakeLLM())
    run, t = h.run_async("Update https://onedrive.live.com/edit.aspx?id=abc mark low attendance")
    pending = wait_pending(run, "source")
    assert "OneDrive" in pending["problem"]
    run.cancel()
    t.join(10)
    assert run.status == RunStatus.CANCELLED


def test_clarifying_question_then_answer(make_harness):
    llm = FakeLLM({"interpret": [goal(clarify="Overall attendance or per subject?"), goal()],
                   "generate_sql": [proposal("SELECT roll_number FROM students")]})
    h = make_harness(llm)
    run, t = h.run_async("Low attendance students", h.source())
    pending = wait_pending(run, "clarification")
    assert pending["question"].startswith("Overall")
    run.respond("clarification", {"answer": "Overall"})
    t.join(10)
    assert run.status == RunStatus.COMPLETED
    assert "A: Overall" in llm.prompts("interpret")[1]


def test_sql_review_approval_with_edit(make_harness):
    h = make_harness(FakeLLM({"interpret": [goal()], "generate_sql": [proposal("SELECT roll_number FROM students")]}))
    run, t = h.run_async("roll numbers", h.source(), require_approval=True)
    wait_pending(run, "sql_approval")
    run.respond("sql_approval", {"approved": True, "sql": "DROP TABLE students"})
    wait_until(lambda: (run.pending or {}).get("error"))
    assert "rejected" in run.pending["error"]
    run.respond("sql_approval", {"approved": True, "sql": "SELECT roll_number FROM students WHERE course_id = 1"})
    t.join(10)
    assert run.status == RunStatus.COMPLETED and len(run.rows) == 30


def test_missing_new_column_declaration_gets_helpful_feedback(make_harness, tmp_path):
    xlsx = make_student_workbook(tmp_path / "class.xlsx")
    llm = FakeLLM({"interpret": [goal(operation="modify")],
                   "generate_sql": [proposal(MARK_SQL, {"limit": 0.5}),  # forgot new_columns
                                    proposal(MARK_SQL, {"limit": 0.5}, NEW_STATUS)]})
    h = make_harness(llm)
    run, t = h.run_async("Add a Status column", h.registry.add_local_path(str(xlsx)).id)
    wait_pending(run, "change_approval")
    assert "new_columns" in run.attempts[0]["feedback"]
    assert "new_columns" in llm.prompts("generate_sql")[1]
    run.cancel()
    t.join(10)
